"""Read-only parser for the SG rate worksheet."""
from dataclasses import dataclass, field
from datetime import date, datetime
import re
from pathlib import Path
from openpyxl import load_workbook
from config import SHEET_NAME, S_EVENT_SHEET_NAME


class ParseError(ValueError):
    pass


@dataclass
class ParsedWorkbook:
    general_rates: list = field(default_factory=list)
    brands: list = field(default_factory=list)
    s_events: list = field(default_factory=list)
    warnings: list = field(default_factory=list)
    source_month: str = ""
    latest_effective_date: str = ""


def parse_rate(value):
    if value is None or value == "":
        return None
    if isinstance(value, str):
        value = value.strip().replace(",", "")
        if not value:
            return None
        if value.endswith("%"):
            return float(value[:-1]) / 100
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number / 100 if number > 1 else number


def parse_date(value, year, month=None):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None:
        return None
    s = str(value).strip().replace("~", "").strip()
    m = re.fullmatch(r"(20\d{2})[-./](\d{1,2})[-./](\d{1,2})", s)
    if m:
        try:
            return date(*map(int, m.groups()))
        except ValueError:
            return None
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})", s)
    if m:
        mo, day = map(int, m.groups())
        # A December/January crossing can appear in one table.
        yr = year + (1 if month == 12 and mo == 1 else -1 if month == 1 and mo == 12 else 0)
        try:
            return date(yr, mo, day)
        except ValueError:
            return None
    return None


def _text(value):
    return str(value).strip() if value is not None else ""


def _parse_s_events(wb, year, month):
    if S_EVENT_SHEET_NAME not in wb.sheetnames:
        raise ParseError(f"无法识别S活动表。请确认是否包含“{S_EVENT_SHEET_NAME}”。")
    rows = wb[S_EVENT_SHEET_NAME].iter_rows(values_only=True)
    header = None
    for row in rows:
        labels = [_text(v).upper() for v in row]
        if '品牌' in labels and 'AGING' in labels and any(v.startswith('REF NO') for v in labels):
            header = row
            break
    if header is None:
        raise ParseError('无法识别S活动表头。')
    labels = [_text(v).upper() for v in header]
    brand_col, aging_col = labels.index('品牌'), labels.index('AGING')
    ref_col = next(i for i, v in enumerate(labels) if v.startswith('REF NO'))
    dates = []
    seen_dates = set()
    for j in range(ref_col + 1, len(header)):
        d = parse_date(header[j], year, month)
        if d and d.isoformat() not in seen_dates:
            dates.append((j, d))
            seen_dates.add(d.isoformat())
    if not dates:
        raise ParseError('无法识别S活动适用日期。')
    events = []
    for row in rows:
        brand = _text(row[brand_col]) if brand_col < len(row) else ''
        aging = _text(row[aging_col]) if aging_col < len(row) else ''
        ref = _text(row[ref_col]) if ref_col < len(row) else ''
        if not brand or not ref:
            continue
        rates = [dict(rate=rate, effective_date=d.isoformat()) for j, d in dates
                 if j < len(row) and (rate := parse_rate(row[j])) is not None]
        if rates:
            events.append(dict(brand_name=brand, aging=aging, ref_no=ref, rates=rates))
    if not events:
        raise ParseError('未找到有效的S活动点数。')
    return events


def parse_excel(source, filename=""):
    try:
        wb = load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        raise ParseError("无法读取 Excel 文件。") from exc
    try:
        if SHEET_NAME not in wb.sheetnames:
            raise ParseError("无法识别该SG点数表格式。请确认是否包含“①SG点数表”。")
        ws = wb[SHEET_NAME]
        rows = list(ws.iter_rows(values_only=True))
        title = " ".join(_text(v) for row in rows[:8] for v in row if v is not None)
        match = re.search(r"(?:^|\D)(\d{2,4})[.\-/](\d{1,2})\s*月?\s*SG", title, re.I)
        if not match:
            match = re.search(r"(\d{2,4})[.\-/](\d{1,2})\s*月", title)
        if not match:
            raise ParseError("无法识别表格年份和月份。")
        year = int(match.group(1))
        year += 2000 if year < 100 else 0
        month = int(match.group(2))
        if not 1 <= month <= 12:
            raise ParseError("表格月份无效。")
        parsed = ParsedWorkbook(source_month=f"{year:04d}-{month:02d}")
        header_idx = None
        for i, row in enumerate(rows):
            vals = [_text(v) for v in row]
            if "支付方式" in vals and "分类" in vals and "品牌" in vals:
                header_idx = i
                break
        if header_idx is None:
            raise ParseError("无法识别品牌表头。")
        header = rows[header_idx]
        labels = [_text(v) for v in header]
        pay_col, cat_col, brand_col = (labels.index(x) for x in ("支付方式", "分类", "品牌"))
        note_col = labels.index("备注") if "备注" in labels else None
        dates = [(j, d) for j, v in enumerate(header) if j > brand_col and (d := parse_date(v, year, month))]
        if not dates:
            raise ParseError("无法识别品牌适用日期。")
        parsed.latest_effective_date = max(d for _, d in dates).isoformat()
        # General-rate header occurs above the brand section.
        general_header = None
        for i, row in enumerate(rows[:header_idx]):
            vals = [_text(v) for v in row]
            if "支付方式" in vals and "分类" in vals:
                general_header = (i, vals.index("支付方式"), vals.index("分类"),
                                  [(j, d) for j, v in enumerate(row) if j > vals.index("分类") and (d := parse_date(v, year, month))])
        if general_header is None or not general_header[3]:
            raise ParseError("无法识别一般点表头或适用日期。")
        gi, gp, gc, gd = general_header
        for row in rows[gi + 1:header_idx]:
            category = _text(row[gc]) if gc < len(row) else ""
            if "一般点" not in category:
                continue
            for j, d in gd:
                rate = parse_rate(row[j]) if j < len(row) else None
                if rate is not None:
                    parsed.general_rates.append(dict(category=category, payment_method=_text(row[gp]), rate=rate, effective_date=d.isoformat()))
        if not parsed.general_rates:
            raise ParseError("未找到有效的一般点数据。")
        seen = {}
        for row in rows[header_idx + 1:]:
            name = _text(row[brand_col]) if brand_col < len(row) else ""
            if not name:
                continue
            rates = [dict(rate=rate, effective_date=d.isoformat()) for j, d in dates
                     if j < len(row) and (rate := parse_rate(row[j])) is not None]
            if not rates:
                parsed.warnings.append(f"{name}: 无有效点数")
                continue
            brand = dict(brand_name=name, category=_text(row[cat_col]), payment_method=_text(row[pay_col]),
                         note=_text(row[note_col]) if note_col is not None and note_col < len(row) else "", rates=rates)
            parsed.brands.append(brand)
            key = " ".join(name.upper().split())
            seen[key] = seen.get(key, 0) + 1
        parsed.warnings.extend(f"重复品牌: {name} ({count})" for name, count in seen.items() if count > 1)
        if not parsed.brands:
            raise ParseError("未找到有效的个别品牌点数。")
        parsed.s_events = _parse_s_events(wb, year, month)
        return parsed
    finally:
        wb.close()

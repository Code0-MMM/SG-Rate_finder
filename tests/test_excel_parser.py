import unittest
from io import BytesIO
from openpyxl import Workbook
from excel_parser import parse_excel, ParseError


def sample():
    wb = Workbook()
    ws = wb.active
    ws.title = '①SG点数表'
    ws['A1'] = '□ 26.09月 SG点数表'
    ws['B7'], ws['C7'], ws['E7'] = '支付方式', '分类', '9/1~'
    for row, category, rate in [(8, '化妆品外其他 一般点', .08), (9, 'LTSC 一般点', .13),
                                (10, '本土C&P 一般点', .18), (11, '进口C&P 一般点', .18)]:
        ws.cell(row, 2, '券扣4%')
        ws.cell(row, 3, category)
        ws.cell(row, 5, rate)
    for col, value in {'B':'支付方式','C':'分类','D':'品牌','E':'9/25~','F':'9/20~','G':'9/12~','N':'备注'}.items():
        ws[f'{col}14'] = value
    for row, name, rate, note in [(15,'BURBERRY',.33,''),(16,'CHANEL',.10,''),(17,'HERMES',.17,'LLG 0%'),(18,'ZERO',0,'')]:
        ws.cell(row, 2, '券扣4%')
        ws.cell(row, 3, 'LF')
        ws.cell(row, 4, name)
        ws.cell(row, 7, rate)
        ws.cell(row, 14, note)
    ws['F15'] = .34
    ws['E15'] = .35
    s = wb.create_sheet('③S活动')
    for col, value in {'B':'品牌','C':'AGING','D':'REF NO.','E':'9/25~','F':'9/1~'}.items():
        s[f'{col}3'] = value
    for col, value in {'B':'AMI','C':'240911','D':'AM3651CL13_3','E':.29,'F':.27}.items():
        s[f'{col}4'] = value
    s['B5'], s['C5'], s['D5'], s['E5'] = 'TUMI', 241111, '1192259RTJ2', 0
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class ParserTests(unittest.TestCase):
    def test_parse(self):
        p = parse_excel(sample())
        self.assertEqual(len(p.general_rates), 4)
        self.assertEqual(len(p.brands), 4)
        self.assertEqual(p.brands[2]['note'], 'LLG 0%')
        self.assertEqual(p.brands[3]['rates'][0]['rate'], 0)
        self.assertEqual(len(p.s_events), 2)
        self.assertEqual(p.s_events[0]['ref_no'], 'AM3651CL13_3')
        self.assertEqual(p.s_events[1]['aging'], '241111')
        self.assertEqual(p.s_events[1]['rates'][0]['rate'], 0)

    def test_bad_workbook(self):
        buf = BytesIO()
        Workbook().save(buf)
        buf.seek(0)
        with self.assertRaises(ParseError):
            parse_excel(buf)
